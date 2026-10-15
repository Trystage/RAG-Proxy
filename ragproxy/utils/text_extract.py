"""从常见文档格式中提取纯文本：Markdown / TXT / PDF / HTML / Word。

不引入任何本地模型推理依赖。
"""
from __future__ import annotations

import re
from pathlib import Path

from bs4 import BeautifulSoup


class TextExtractError(RuntimeError):
    pass


def extract_markdown(path: Path) -> str:
    return path.read_text(encoding="utf-8", errors="replace")


def extract_txt(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            return raw.decode(enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return raw.decode("utf-8", errors="replace")


def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:  # pragma: no cover
        raise TextExtractError("未安装 pypdf，无法解析 PDF") from e
    reader = PdfReader(str(path))
    pages = []
    for page in reader.pages:
        try:
            pages.append(page.extract_text() or "")
        except Exception:  # noqa: BLE001 - 单页失败不阻断
            continue
    return "\n\n".join(pages)


def extract_html(path: Path) -> str:
    raw = path.read_bytes()
    for enc in ("utf-8", "utf-8-sig", "gb18030", "latin-1"):
        try:
            html = raw.decode(enc)
            break
        except (UnicodeDecodeError, LookupError):
            continue
    else:  # pragma: no cover
        html = raw.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    for tag in soup(["script", "style", "noscript", "template", "svg"]):
        tag.decompose()
    text = soup.get_text("\n", strip=True)
    return re.sub(r"\n{3,}", "\n\n", text)


def extract_docx(path: Path) -> str:
    try:
        import docx  # python-docx
    except ImportError as e:  # pragma: no cover
        raise TextExtractError("未安装 python-docx，无法解析 Word 文档") from e
    document = docx.Document(str(path))
    parts: list[str] = []
    for para in document.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_doc(path: Path) -> str:
    """旧版 .doc（二进制）无法用纯 python-docx 解析，给出明确提示。"""
    raise TextExtractError(
        "旧版 .doc 二进制格式暂不支持解析，请转换为 .docx / .txt / .md 后上传"
    )


_EXTRACTORS = {
    ".md": extract_markdown,
    ".markdown": extract_markdown,
    ".txt": extract_txt,
    ".pdf": extract_pdf,
    ".html": extract_html,
    ".htm": extract_html,
    ".docx": extract_docx,
    ".doc": extract_doc,
}


def extract_text(path: Path) -> str:
    ext = path.suffix.lower()
    fn = _EXTRACTORS.get(ext)
    if fn is None:
        raise TextExtractError(f"不支持的文件类型：{ext}")
    return fn(path)
